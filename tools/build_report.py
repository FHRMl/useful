from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "报告.docx"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x20, 0x20, 0x20)
MUTED = RGBColor(0x66, 0x66, 0x66)
HEADER_FILL = "F2F4F7"
BORDER = "D7DDE5"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_p.paragraph_format.space_after = Pt(0)
    left = header_p.add_run("信息组织大作业")
    set_run_font(left, size=9, color=MUTED, bold=True)
    header_p.add_run("    ")
    right = header_p.add_run("课程知识图谱社区系统")
    set_run_font(right, size=9, color=MUTED)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_after = Pt(0)
    run = footer_p.add_run("课程知识图谱社区系统报告")
    set_run_font(run, size=9, color=MUTED)


def paragraph_bottom_border(paragraph, color="A6B4C4", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("课程知识图谱社区系统报告")
    set_run_font(run, size=23, color=INK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("信息组织大作业 2026")
    set_run_font(run, size=13, color=MUTED)

    metadata = [
        ("课程", "信息组织"),
        ("系统", "课程知识图谱社区系统"),
        ("技术栈", "Flask / MySQL / Neo4j / ECharts"),
        ("日期", "2026-07-02"),
        ("小组分工", "待补充：成员姓名、学号与具体职责"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}：")
        set_run_font(label_run, size=11, color=INK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=11, color=INK)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(12)
    paragraph_bottom_border(rule)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_borders(table)
    set_table_width(table, widths)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, HEADER_FILL)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=10.5, color=INK, bold=True)

    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            cell = row.cells[index]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=10.5, color=INK)

    set_table_width(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(4)
    return table


def add_para(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=11, color=INK)


def build():
    doc = Document()
    configure_document(doc)
    add_title_block(doc)

    doc.add_heading("一、系统概述", level=1)
    add_para(
        doc,
        "本系统以信息管理学院课程为核心对象，将课程、知识点、中图法分类号和四类分类维度组织为知识图谱，并在 Web 端提供课程社区功能。系统把稳定的图谱结构存入 Neo4j，把用户、课程评价和知识点标注等运行时数据存入 MySQL，实现信息组织理论与社区交互功能的结合。",
    )

    doc.add_heading("二、作业要求对应情况", level=1)
    add_table(
        doc,
        ["作业要求", "实现情况", "对应文件或页面"],
        [
            ("概念模型/本体", "定义课程、知识点、中图法分类、四类维度和用户动态数据类。", "ontology/course_ontology.owl；docs/data_design.md"),
            ("Neo4j 图数据库", "提供 6 门课程、30 个知识点、分类号、维度和主题词关系。", "neo4j/init.cypher"),
            ("MySQL 关系数据库", "保存用户、课程评分点评、知识点标注，并初始化 3 个虚拟用户。", "mysql/schema.sql"),
            ("多用户功能", "支持注册、登录、注销、课程评分点评和知识点标注。", "backend/app.py；课程详情页"),
            ("多页面 Web", "包含主页、注册/登录、个人主页、课程详情、知识图谱、分析页。", "frontend/templates/"),
            ("分析展示", "展示度数最高知识点、共享知识点最多课程、评分排行和维度分布。", "/analysis；/api/analysis"),
        ],
        [1.45, 3.0, 2.05],
    )

    doc.add_heading("三、本体与数据模型", level=1)
    add_para(
        doc,
        "课程图谱采用 Course、KnowledgePoint、CLCClass、DimensionValue 四类核心图谱节点。每门课程通过 HAS_KNOWLEDGE_POINT 连接至少 5 个知识点，通过 CLASSIFIED_BY 连接中图法分类号，并通过 HAS_DIMENSION 分别连接文化、机构、个人、计算四类维度取值。",
    )
    add_table(
        doc,
        ["数据对象", "数据库", "设计理由"],
        [
            ("课程、知识点、分类号、维度", "Neo4j", "结构稳定且关系密集，适合多跳查询、共享知识点分析和力导向图展示。"),
            ("用户账号", "MySQL", "需要唯一约束、登录认证和事务性更新。"),
            ("课程评分点评", "MySQL", "属于用户生成的动态数据，需要按用户和课程维护当前评价。"),
            ("知识点标注", "MySQL", "属于个体学习记录，可区分公开和私有标注。"),
        ],
        [1.45, 1.05, 4.0],
    )

    doc.add_heading("四、数据来源与初始化", level=1)
    add_para(
        doc,
        "课程样例参考南京大学本科生院与信息管理学院公开课表，选择信息组织、信息检索、数据结构、数据科学与数据分析、管理信息系统、信息分析六门课程。知识点按照《汉语主题词表》主题词粒度整理，并为课程添加参考性中图法分类号。",
    )
    add_bullets(
        doc,
        [
            "Neo4j 初始化脚本使用 MERGE 写入课程、知识点、分类号、维度值和知识点关系，便于重复执行。",
            "MySQL 初始化脚本创建 3 个虚拟用户，每个用户至少评价 3 门课程，并写入若干知识点标注。",
            "Flask 另提供 SQLite + 本地图谱 fallback，用于没有数据库服务时快速预览界面；正式运行仍推荐使用 MySQL + Neo4j。",
        ],
    )

    doc.add_heading("五、功能模块说明", level=1)
    add_table(
        doc,
        ["模块", "页面/接口", "说明"],
        [
            ("课程主页", "/", "展示课程列表、学分、分类号、知识点数量和平均评分。"),
            ("用户认证", "/register、/login、/logout", "完成注册、登录、退出和会话管理。"),
            ("课程详情", "/courses/<course_id>", "展示课程知识点、维度和点评；登录后可评分和标注。"),
            ("个人主页", "/me", "汇总当前用户的课程点评和知识点标注。"),
            ("知识图谱", "/graph、/api/graph", "ECharts 力导向图展示课程、知识点、分类和维度关系。"),
            ("分析展示", "/analysis、/api/analysis", "统计知识点度数、课程共享知识点、评分排行和维度分布。"),
        ],
        [1.25, 1.85, 3.4],
    )

    doc.add_heading("六、信息组织理论应用", level=1)
    add_table(
        doc,
        ["理论/方法", "系统中的应用"],
        [
            ("分类法", "为课程添加中图法分类号，使课程可按学科知识体系定位。"),
            ("主题词表", "知识点按主题词粒度统一命名，减少同义表达导致的检索和统计偏差。"),
            ("本体建模", "用类、属性和关系约束课程图谱结构，形成可复用的概念模型。"),
            ("知识关联", "用上下位、相关、组成、应用等关系表达知识点之间的语义联系。"),
            ("用户标注", "把个人理解作为运行时数据保存，支持课程社区中的知识再组织。"),
        ],
        [1.55, 4.95],
    )

    doc.add_heading("七、运行与复现", level=1)
    add_para(
        doc,
        "完整运行时先通过 docker compose 启动 MySQL 和 Neo4j，再导入 neo4j/init.cypher，随后安装 Python 依赖并启动 Flask。README.md 已给出 PowerShell 命令、环境变量和演示用户。若只需检查界面，可直接使用 SQLite 与本地图谱 fallback。",
    )

    doc.add_heading("八、未来扩展设想", level=1)
    add_bullets(
        doc,
        [
            "增加课程大纲、教材、授课教师等实体，扩展课程知识图谱深度。",
            "基于用户评分和标注构建个性化课程推荐。",
            "加入图谱路径查询，例如从某一知识点反查相关课程和先修知识。",
            "将《汉语主题词表》和中图法的更多标准数据导入为可复用权威词表节点。",
        ],
    )

    doc.add_heading("九、小组分工", level=1)
    add_para(
        doc,
        "待补充：提交前请填写成员姓名、学号和具体分工，例如数据整理、本体设计、后端开发、前端可视化、报告撰写与测试。",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)

